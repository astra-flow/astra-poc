"""
重新计算：PRD生成费用 580→120 后的预算
"""
import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('敏捷软件开发AI应用需求情况统计.xlsx', data_only=True)
ws1 = wb['应用产品']; ws2 = wb['基础软件']

MARKET_WP = {
    'DeepSeek-V4': 18.00,
    'Qwen3-Plus': 4.20,
    'GLM-5.2': 18.00,
    'Codex': 37.80,
    'Doubao-Seed-Code': 4.20,
    'Gemini': 21.60,
}

def norm(raw):
    if not raw or raw in ('无','-','—','None'): return None
    s=str(raw).lower().replace(' ','').replace('\n','').replace('_','').replace('，','')
    if 'deepseek' in s: return 'DeepSeek-V4'
    if 'qwen' in s or '千问' in s: return 'Qwen3-Plus'
    if 'glm' in s or 'gml' in s: return 'GLM-5.2'
    if 'doubao' in s or 'duubao' in s: return 'Doubao-Seed-Code'
    if 'gemini' in s: return 'Gemini'
    if 'codex' in s: return 'Codex'
    return None

scenes_app = [('PRD生成',9,10,11,12,13),('架构图生成',16,17,18,19,20),('代码生成',23,24,25,26,27),('测试脚本',30,31,32,33,34)]

agg = defaultdict(lambda: {'月总费用':0,'记录数':0,'场景集':set()})
scene_agg = defaultdict(lambda: {'fee':0,'persons':0})

# 应用产品
for r in range(3, ws1.max_row+1):
    p = ws1.cell(row=r,column=2).value
    if not p: continue
    for scene, mcol, tcol, fcol, pcol, total_col in scenes_app:
        raw = ws1.cell(row=r, column=mcol).value
        model = norm(raw)
        if not model: continue
        fee = ws1.cell(row=r, column=fcol).value or 0
        persons = ws1.cell(row=r, column=pcol).value or 0
        total = ws1.cell(row=r, column=total_col).value or 0
        try: fee=float(fee); persons=float(persons); total=float(total)
        except: pass
        agg[model]['月总费用'] += total
        agg[model]['记录数'] += 1
        agg[model]['场景集'].add(scene)
        scene_agg[scene]['fee'] += total
        scene_agg[scene]['persons'] += persons

# 基础软件
for r in range(3, ws2.max_row+1):
    m = ws2.cell(row=r,column=1).value
    if not m: continue
    fd=float(ws2.cell(row=r,column=7).value or 0); pd=float(ws2.cell(row=r,column=6).value or 0)
    fl=float(ws2.cell(row=r,column=14).value or 0); pl=float(ws2.cell(row=r,column=13).value or 0)
    for sm in ['DeepSeek-V4','Qwen3-Plus','GLM-5.2']:
        agg[sm]['月总费用'] += fd/3 + fl/3
        agg[sm]['记录数'] += 2
        agg[sm]['场景集'].add('基础软件开发')
        agg[sm]['场景集'].add('日志分析')
    scene_agg['基础软件开发']['fee'] += fd
    scene_agg['基础软件开发']['persons'] += pd
    scene_agg['日志分析']['fee'] += fl
    scene_agg['日志分析']['persons'] += pl

print('='*100)
print('一、按模型族聚合（更新后）')
print('='*100)
print(f'{"模型族":<16}{"月费用(元)":>12}{"记录数":>8}{"加权单价":>10}{"月Token(M)":>14}{"年Token(万)":>14}{"波动1.15(万)":>16}{"年预算(万)":>12}')
print('-'*100)
total_month = 0; total_year = 0
rows = []
for m in sorted(agg.keys(), key=lambda x:-agg[x]['月总费用']):
    d = agg[m]
    wp = MARKET_WP[m]
    mt = d['月总费用'] / wp  # 百万
    yt = mt * 12 * 100  # 万
    yb = yt * 1.15
    yf = d['月总费用'] * 12 * 1.15 / 10000
    print(f'{m:<16}{d["月总费用"]:>12.0f}{d["记录数"]:>8}{wp:>10.2f}{mt:>14.2f}{yt:>14.0f}{yb:>16.0f}{yf:>12.1f}')
    total_month += d['月总费用']; total_year += yf
    rows.append((m, d['月总费用'], wp, mt, yt, yb, yf))
print('-'*100)
print(f'{"合计":<16}{total_month:>12.0f}{"":>8}{"":>10}{"":>14}{"":>14}{"":>16}{total_year:>12.1f}')

print()
print('='*100)
print('二、按场景聚合（更新后）')
print('='*100)
print(f'{"场景":<16}{"月费用(元)":>14}{"使用人数":>10}{"人均月费":>10}{"占比":>8}')
print('-'*60)
scene_total = 0
for s in sorted(scene_agg.keys(), key=lambda x:-scene_agg[x]['fee']):
    d = scene_agg[s]
    pct = d['fee']/total_month*100
    pc = d['fee']/d['persons'] if d['persons']>0 else 0
    print(f'{s:<16}{d["fee"]:>14.0f}{d["persons"]:>10.0f}{pc:>10.0f}{pct:>7.1f}%')
    scene_total += d['fee']
print('-'*60)

print()
print('='*100)
print('三、人均费用（400人口径）')
print('='*100)
print(f'月总费用: {total_month:.0f} 元')
print(f'人均月费用: {total_month/400:.0f} 元/人/月')
print(f'人均年费用: {total_month/400*12:.0f} 元/人/年')
print(f'人均年预算(×1.15): {total_month/400*12*1.15:.0f} 元/人/年')
print(f'400人年预算: {total_month/400*12*1.15*400/10000:.1f} 万元')

print()
print('='*100)
print('四、与更新前对比')
print('='*100)
old_total = 250720
old_year = 346.0
print(f'{"指标":<20}{"更新前":>14}{"更新后":>14}{"差额":>14}{"降幅":>10}')
print('-'*72)
print(f'{"月总费用(元)":<20}{old_total:>14.0f}{total_month:>14.0f}{total_month-old_total:>14.0f}{(total_month-old_total)/old_total*100:>9.1f}%')
print(f'{"年预算(万元)":<20}{old_year:>14.1f}{total_year:>14.1f}{total_year-old_year:>14.1f}{(total_year-old_year)/old_year*100:>9.1f}%')

# PRD生成场景对比
print()
print('='*100)
print('五、PRD生成场景对比（主要变化点）')
print('='*100)
print(f'{"指标":<24}{"更新前":>14}{"更新后":>14}')
print('-'*56)
# 更新前 PRD 费用 26680，更新后重新算
prd_new = scene_agg['PRD生成']['fee']
print(f'{"PRD生成月费用(元)":<24}{26680:>14.0f}{prd_new:>14.0f}')
print(f'{"PRD生成人均月费(元)":<24}{580:>14.0f}{120:>14.0f}')
