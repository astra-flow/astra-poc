"""
=====================================================================
供应商询价 Excel 模板生成脚本
=====================================================================
数据来源：data_provider.py（统一数据源）
功能：生成标准化的供应商报价 Excel 模板

使用方式：
  python3 scripts/04_generate_inquiry_template.py

输出：
  output/Token资源包报价模板.xlsx
=====================================================================
"""
import os
import sys
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 添加 scripts 目录到 path
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPTS_DIR)

from data_provider import get_all_data

# ============ 路径 ============
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(BASE_DIR, 'output')

# ============ 样式 ============
title_font = Font(name='微软雅黑', size=14, bold=True, color='FFFFFF')
title_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
section_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
section_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
required_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')  # 黄色=必填
thin_border = Border(
    left=Side(style='thin', color='999999'),
    right=Side(style='thin', color='999999'),
    top=Side(style='thin', color='999999'),
    bottom=Side(style='thin', color='999999')
)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)


def set_cell(ws, row, col, value, font=None, fill=None, alignment=None, border=None):
    """设置单元格"""
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border
    return cell


def generate_inquiry_template():
    """生成供应商询价 Excel 模板"""
    data = get_all_data()

    wb = openpyxl.Workbook()

    # ============ Sheet 1: Token资源包报价 ============
    ws = wb.active
    ws.title = 'Token资源包报价'

    # 标题
    ws.merge_cells('A1:K1')
    set_cell(ws, 1, 1, 'Token 资源包报价表', title_font, title_fill, center)
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:K2')
    set_cell(ws, 2, 1, '供应商填写（黄色单元格为必填项）', Font(name='微软雅黑', size=10, italic=True), None, left)

    # 供应商信息
    set_cell(ws, 3, 1, '供应商名称', header_font, header_fill, center, thin_border)
    set_cell(ws, 3, 2, '', None, required_fill, center, thin_border)
    ws.merge_cells('B3:D3')
    set_cell(ws, 3, 5, '联系人', header_font, header_fill, center, thin_border)
    set_cell(ws, 3, 6, '', None, required_fill, center, thin_border)
    set_cell(ws, 3, 7, '联系方式', header_font, header_fill, center, thin_border)
    set_cell(ws, 3, 8, '', None, required_fill, center, thin_border)
    set_cell(ws, 3, 9, '报价日期', header_font, header_fill, center, thin_border)
    set_cell(ws, 3, 10, '', None, required_fill, center, thin_border)

    # 一、采购需求概况
    ws.merge_cells('A5:K5')
    set_cell(ws, 5, 1, '一、采购需求概况（采购方提供，供应商确认）', section_font, section_fill, left)
    ws.row_dimensions[5].height = 25

    requirements = [
        ('使用人数', f'约 {data["PERSONS"]} 人（研发团队）'),
        ('使用场景', '代码生成、PRD生成、架构图生成、测试脚本、基础软件开发、日志分析'),
        ('年度Token需求量', '约 1-2 万亿 Token（规模参考，实际按框架协议签约）'),
        ('主力模型要求', '必须包含：DeepSeek V4 (Pro)、DeepSeek V4 Flash、GLM 5.2'),
        ('计价单位', '统一按 百万Token 报价（元/百万Token），不接受 credit/积分等非标准计价方式'),
    ]
    for i, (key, val) in enumerate(requirements):
        r = 6 + i
        set_cell(ws, r, 1, '', None, None, center, thin_border)
        set_cell(ws, r, 2, key, Font(name='微软雅黑', size=10, bold=True), PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid'), center, thin_border)
        ws.merge_cells(f'C{r}:K{r}')
        set_cell(ws, r, 3, val, Font(name='微软雅黑', size=10), None, left, thin_border)

    # 二、模型单价报价
    ws.merge_cells('A12:K12')
    set_cell(ws, 12, 1, '二、模型单价报价（供应商填写）', section_font, section_fill, left)
    ws.row_dimensions[12].height = 25

    headers_model = ['序号', '模型名称', '模型档位\n(旗舰/轻量/代码)', '输入单价\n(缓存命中)\n元/百万Token',
                     '输入单价\n(缓存未命中)\n元/百万Token', '输出单价\n元/百万Token',
                     '资源包类型\n(套餐/阶梯/按量)', '资源包规格\n(如10亿Token包)', '资源包价格\n(元)',
                     '折扣力度\n(对比刊例价)', '备注']
    for i, h in enumerate(headers_model):
        set_cell(ws, 13, i + 1, h, header_font, header_fill, center, thin_border)
    ws.row_dimensions[13].height = 45

    # 必填模型（前3行）
    required_models = [
        ('DeepSeek V4 (Pro)', '旗舰'),
        ('DeepSeek V4 Flash', '轻量'),
        ('GLM 5.2', '旗舰'),
    ]
    for i, (name, tier) in enumerate(required_models):
        r = 14 + i
        set_cell(ws, r, 1, i + 1, Font(name='微软雅黑', size=10), None, center, thin_border)
        set_cell(ws, r, 2, name, Font(name='微软雅黑', size=10), None, left, thin_border)
        set_cell(ws, r, 3, tier, Font(name='微软雅黑', size=10), None, center, thin_border)
        for c in range(4, 12):
            set_cell(ws, r, c, '', None, required_fill, center, thin_border)

    # 空行（供应商自填）
    for i in range(3, 7):
        r = 14 + i
        set_cell(ws, r, 1, i + 1, Font(name='微软雅黑', size=10), None, center, thin_border)
        for c in range(2, 12):
            set_cell(ws, r, c, '', None, None, center, thin_border)

    # 三、年度总价测算
    ws.merge_cells('A21:K21')
    set_cell(ws, 21, 1, '三、年度总价测算（供应商填写）', section_font, section_fill, left)
    ws.row_dimensions[21].height = 25

    headers_plan = ['方案', '报价方案说明', '年度总价\n(万元，含税)',
                    '是否支持框架协议\n(按金额签约/按量结算)', '套餐有效期', '是否支持结转',
                    '超额后付费费率\n(元/百万Token)', '备注']
    for i, h in enumerate(headers_plan):
        set_cell(ws, 22, i + 1, h, header_font, header_fill, center, thin_border)
    ws.row_dimensions[22].height = 45

    for i in range(3):
        r = 23 + i
        set_cell(ws, r, 1, f'方案{i + 1}', Font(name='微软雅黑', size=10), None, center, thin_border)
        for c in range(2, 9):
            set_cell(ws, r, c, '', None, required_fill, center, thin_border)

    # 四、Credit换算确认
    ws.merge_cells('A27:K27')
    set_cell(ws, 27, 1, '四、Credit换算确认（如底层采用Credit计费，必填）', section_font, section_fill, left)
    ws.row_dimensions[27].height = 25

    ws.merge_cells('A28:K28')
    set_cell(ws, 28, 1, '各供应商的credits换算规则在官方平台有公开说明，此处仅需确认精确值并引用官方文档。',
             Font(name='微软雅黑', size=9, italic=True, color='666666'), None, left)

    headers_credit = ['序号', 'Credit定义', '1 Credit = 多少 Token\n(精确值，非区间)',
                      '1 元 = 多少 Credit', '官方文档链接']
    for i, h in enumerate(headers_credit):
        set_cell(ws, 29, i + 1, h, header_font, header_fill, center, thin_border)
    ws.row_dimensions[29].height = 40

    for i in range(3):
        r = 30 + i
        set_cell(ws, r, 1, i + 1, Font(name='微软雅黑', size=10), None, center, thin_border)
        for c in range(2, 6):
            set_cell(ws, r, c, '', None, required_fill, center, thin_border)

    # 说明
    ws.merge_cells('A34:K34')
    set_cell(ws, 34, 1, '说明：1. 所有单价单位为 元/百万Token；2. 黄色单元格为供应商必填项；3. credits换算规则引用官方文档即可。',
             Font(name='微软雅黑', size=9, color='666666'), None, left)

    # 列宽
    col_widths = [6, 22, 14, 14, 16, 12, 14, 16, 14, 14, 30]
    for i, w in enumerate(col_widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    # ============ Sheet 2: 填写说明 ============
    ws2 = wb.create_sheet('填写说明')
    ws2.column_dimensions['A'].width = 4
    ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 70

    set_cell(ws2, 1, 1, '填写说明', title_font, title_fill, center)
    ws2.merge_cells('A1:C1')
    ws2.row_dimensions[1].height = 30

    instructions = [
        ('背景', '我司正在进行企业级AI平台Token资源包采购，需统一口径横向比价。'),
        ('计价单位', '统一按 百万Token 报价（元/百万Token），不接受 credit/积分等非标准计价方式。'),
        ('必填模型', 'DeepSeek V4 (Pro)、DeepSeek V4 Flash、GLM 5.2 三款为必填。'),
        ('缓存命中', '输入单价请区分"缓存命中"和"缓存未命中"两档，DeepSeek官方即如此计价。'),
        ('资源包类型', '请分别提供套餐包、阶梯包、按量包三种形态的价格。'),
        ('年度总价', '按 1-2 万亿 Token 年度需求测算总价（万元，含税），便于横向比价。'),
        ('框架协议', '是否支持"按年度金额签约、按实际消耗结算"的框架协议模式，请明确。'),
        ('Credit换算', '各供应商credits换算规则在官方平台有公开说明，第四节仅需确认精确值并引用官方文档链接。'),
        ('替代模型', '如不提供主力模型，需提供替代模型并在备注中说明对比关系（附链接）。'),
        ('时间要求', '请于收到本询价后 3 个工作日内提供完整报价。'),
        ('联系人', '如有疑问请联系采购方。'),
    ]
    for i, (key, val) in enumerate(instructions):
        r = 3 + i
        set_cell(ws2, r, 1, '', None, None, center, thin_border)
        set_cell(ws2, r, 2, key, Font(name='微软雅黑', size=10, bold=True), PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid'), center, thin_border)
        set_cell(ws2, r, 3, val, Font(name='微软雅黑', size=10), None, left, thin_border)
        ws2.row_dimensions[r].height = 30

    # ============ Sheet 3: 参考基准价（采购方内部，不发送） ============
    ws3 = wb.create_sheet('参考基准价(采购方)')
    ws3.column_dimensions['A'].width = 6
    ws3.column_dimensions['B'].width = 22
    ws3.column_dimensions['C'].width = 10
    for c in ['D', 'E', 'F']:
        ws3.column_dimensions[c].width = 16
    ws3.column_dimensions['G'].width = 30

    set_cell(ws3, 1, 1, '参考基准价（采购方内部参考，不发送给供应商）', title_font, title_fill, center)
    ws3.merge_cells('A1:G1')
    ws3.row_dimensions[1].height = 30

    headers_ref = ['序号', '模型', '档位', '输入(缓存命中)\n元/百万Token',
                   '输入(缓存未命中)\n元/百万Token', '输出\n元/百万Token', '数据来源']
    for i, h in enumerate(headers_ref):
        set_cell(ws3, 3, i + 1, h, header_font, header_fill, center, thin_border)
    ws3.row_dimensions[3].height = 40

    for i, p in enumerate(data['official_prices_table']):
        r = 4 + i
        set_cell(ws3, r, 1, i + 1, Font(name='微软雅黑', size=10), None, center, thin_border)
        set_cell(ws3, r, 2, p['model'], Font(name='微软雅黑', size=10), None, left, thin_border)
        set_cell(ws3, r, 3, p['tier'], Font(name='微软雅黑', size=10), None, center, thin_border)
        set_cell(ws3, r, 4, p['input_hit'], Font(name='微软雅黑', size=10), None, center, thin_border)
        set_cell(ws3, r, 5, p['input_miss'], Font(name='微软雅黑', size=10), None, center, thin_border)
        set_cell(ws3, r, 6, p['output'], Font(name='微软雅黑', size=10), None, center, thin_border)
        set_cell(ws3, r, 7, f"{p['vendor']}官方", Font(name='微软雅黑', size=10), None, left, thin_border)

    # 说明
    last_row = 4 + len(data['official_prices_table']) + 1
    ws3.merge_cells(f'A{last_row}:G{last_row}')
    set_cell(ws3, last_row, 1,
             '说明：此表为采购方内部参考基准价，用于验证供应商报价合理性。数据来自各厂商官方定价页。',
             Font(name='微软雅黑', size=9, color='666666'), None, left)

    # ============ 保存 ============
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, 'Token资源包报价模板.xlsx')
    wb.save(output_path)
    print(f'✅ 询价模板已生成: {output_path}')
    print(f'   Sheet1: Token资源包报价（供应商填写）')
    print(f'   Sheet2: 填写说明')
    print(f'   Sheet3: 参考基准价（采购方内部，不发送）')


if __name__ == '__main__':
    generate_inquiry_template()
